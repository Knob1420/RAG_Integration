"""TASK-P01-FULL-PIPELINE: R7/R8 XLSX 解析 → DocumentElement。

WPS 修复 → openpyxl 双载（data_only 求值 + 原式兜底）→ 合并单元格展开 →
按空行切逻辑区域 → 表区（多行表头展平，行记录自带表头）/ 元信息区（key_value）。
返回 (elements, stats)；stats 供 quality.py 单元格守恒检查。
"""
from __future__ import annotations

from pathlib import Path

import openpyxl

from .models import (DocumentElement, ElementType, ExtractionMode, Location,
                     Provenance, RouteId)
from .probe import normalize_zip


def _cell_str(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _grid(ws_val, ws_raw) -> list[list[str | None]]:
    """值网格：data_only 求值优先，缓存缺失保留原式；合并单元格展开（守恒计入）。"""
    mr, mc = ws_val.max_row, ws_val.max_column
    grid = [[None] * mc for _ in range(mr)]
    for r in range(1, mr + 1):
        for c in range(1, mc + 1):
            v = ws_val.cell(r, c).value
            if v is None:
                raw = ws_raw.cell(r, c).value
                if isinstance(raw, str) and raw.startswith("="):
                    v = raw  # 求不到 → 原式
            grid[r - 1][c - 1] = _cell_str(v)
    for m in ws_val.merged_cells.ranges:
        tl = grid[m.min_row - 1][m.min_col - 1]
        if tl is None:
            continue
        for r in range(m.min_row, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                if grid[r - 1][c - 1] is None:
                    grid[r - 1][c - 1] = tl
    return grid


def _regions(grid: list[list[str | None]]) -> list[list[int]]:
    """空行分隔的连续非空行块 → [起始行, 结束行]（0 起）。"""
    out, start = [], None
    for i, row in enumerate(grid):
        if any(v is not None for v in row):
            if start is None:
                start = i
        elif start is not None:
            out.append([start, i - 1])
            start = None
    if start is not None:
        out.append([start, len(grid) - 1])
    return out


def _flatten_multi_header(rows: list[list[str | None]]) -> list[str]:
    """首行有列空缺 → 两行按位拼接展平（组名/字段名），无空缺则原样。"""
    head = rows[0]
    if not any(v is None for v in head):
        return [v or "" for v in head]
    second = rows[1] if len(rows) > 1 else []
    out = []
    for i, v in enumerate(head):
        b = second[i] if i < len(second) else None
        out.append("/".join(x for x in (v, b) if x))
    return out


def parse_xlsx(path: Path, source_id: str, route: RouteId,
               route_reason: str) -> tuple[list[DocumentElement], dict]:
    real = normalize_zip(path)
    wb_val = openpyxl.load_workbook(real, data_only=True)
    wb_raw = openpyxl.load_workbook(real, data_only=False)
    prov = Provenance(parser_name="openpyxl",
                      parser_version=openpyxl.__version__,
                      route=route, extraction_mode=ExtractionMode.XLSX_STRUCTURED,
                      route_reason=route_reason)
    elements: list[DocumentElement] = []
    table_idx = 0
    input_cells = output_cells = 0

    for ws_val in wb_val.worksheets:
        ws_raw = wb_raw[ws_val.title]
        hidden = {r for r, d in ws_val.row_dimensions.items() if d.hidden}
        for im in getattr(ws_val, "_images", []):  # 内嵌图占位（锚点格行号）
            frm = getattr(im.anchor, "_from", None)
            r = frm.row + 1 if frm else 1
            elements.append(DocumentElement(
                source_id=source_id, element_type=ElementType.FIGURE,
                content="[figure]",
                location=Location(source_id, sheet_name=ws_val.title, row_no=r),
                provenance=prov,
                metadata={"size": f"{getattr(im, 'width', '?')}x"
                                  f"{getattr(im, 'height', '?')}"}))
        grid = _grid(ws_val, ws_raw)
        input_cells += sum(v is not None for row in grid for v in row)
        for start, end in _regions(grid):
            block = grid[start:end + 1]
            widths = [sum(v is not None for v in row) for row in block]
            if all(w <= 2 for w in widths):  # 元信息区
                for i, row in enumerate(block):
                    vals = [v for v in row if v is not None]
                    if not vals:
                        continue
                    r = start + i + 1
                    content = vals[0] if len(vals) == 1 else f"{vals[0]}: {vals[1]}"
                    output_cells += len(vals)
                    elements.append(DocumentElement(
                        source_id=source_id, element_type=ElementType.KEY_VALUE,
                        content=content,
                        location=Location(source_id, sheet_name=ws_val.title,
                                          row_no=r),
                        provenance=prov,
                        metadata={"hidden": r in hidden} if r in hidden else {}))
                continue
            # 表区
            # 列裁剪：只留块内有值的列区间（尾部/头部空白列不参与表头判定）
            cols = [c for c in range(len(block[0]))
                    if any(r[c] is not None for r in block)]
            block = [r[cols[0]:cols[-1] + 1] for r in block]
            # 合并标题行：整行同值是合并展开痕迹 → 产单条 key_value，表头取下一行
            if len(block) > 1 and all(block[0]) and len(set(block[0])) == 1:
                r0 = start + 1
                output_cells += len(block[0])  # 展开格数计产出（守恒同口径）
                elements.append(DocumentElement(
                    source_id=source_id, element_type=ElementType.KEY_VALUE,
                    content=block[0][0],
                    location=Location(source_id, sheet_name=ws_val.title,
                                      row_no=r0),
                    provenance=prov, metadata={"merged_title": True}))
                block = block[1:]
                start += 1
                cols = [c for c in range(len(block[0]))
                        if any(r[c] is not None for r in block)]
                block = [r[cols[0]:cols[-1] + 1] for r in block]  # 弹出标题后再裁一次
            table_idx += 1
            header = _flatten_multi_header(block)
            data_start = 2 if any(v is None for v in block[0]) and len(block) > 1 else 1
            header_rows = f"{start + 1}-{start + data_start}"
            # 表头进 metadata：按参与表头的行计非空格（与输入同口径）
            for hr in range(data_start):
                output_cells += sum(v is not None for v in block[hr])
            for i in range(data_start, len(block)):
                row = block[i]
                if not any(v is not None for v in row):
                    continue
                r = start + i + 1
                cells = [v or "" for v in row]
                output_cells += sum(bool(c) for c in cells)
                md = {"header": header, "header_row_range": header_rows}
                if r in hidden:
                    md["hidden"] = True
                elements.append(DocumentElement(
                    source_id=source_id, element_type=ElementType.TABLE_ROW,
                    content=" | ".join(cells),
                    location=Location(source_id, sheet_name=ws_val.title,
                                      table_index=table_idx, row_no=r),
                    provenance=prov, metadata=md))
    wb_val.close()
    wb_raw.close()
    return elements, {"input_cells": input_cells, "output_cells": output_cells}
