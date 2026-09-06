"""parse_docx 合成样例测试（不依赖真实资料）。"""
from pathlib import Path

import docx
import pytest

from rag_integration.models import RouteId
from rag_integration.parse_docx import parse_docx

SID = "SRC-0123456789"


@pytest.fixture()
def sample_docx(tmp_path: Path) -> Path:
    d = docx.Document()
    d.add_heading("总体方案", level=1)
    d.add_paragraph("合成段落内容。")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "参数"
    t.cell(0, 1).text = "值"
    t.cell(1, 0).text = "质量"
    t.cell(1, 1).text = "12kg"
    p = tmp_path / "s.docx"
    d.save(p)
    return p


def test_parse_docx_structure(sample_docx):
    els = parse_docx(sample_docx, SID, RouteId.R4, "standard")
    types = [e.element_type.value for e in els]
    assert types.count("heading") == 1
    assert types.count("paragraph") == 1
    assert types.count("table_row") == 1

    h = next(e for e in els if e.element_type.value == "heading")
    assert h.metadata["level"] == 1
    assert h.location.section_chain == ("总体方案",)

    row = next(e for e in els if e.element_type.value == "table_row")
    assert row.metadata["header"] == ["参数", "值"]
    assert row.content == "质量 | 12kg"
    assert row.location.table_index == 1 and row.location.row_no == 2
    assert row.provenance.extraction_mode.value == "docx_structured"


def test_merged_title_row_becomes_kv(tmp_path):
    d = docx.Document()
    d.add_paragraph("正文")
    t = d.add_table(rows=3, cols=3)
    # 首行横向合并，填同一标题（python-docx 合并后 cells 重复同值）
    t.cell(0, 0).text = "封面标题表"
    t.cell(0, 1).text = "封面标题表"
    t.cell(0, 2).text = "封面标题表"
    t.cell(1, 0).text = "参数"
    t.cell(1, 1).text = "值"
    t.cell(1, 2).text = "备注"
    t.cell(2, 0).text = "质量"
    t.cell(2, 1).text = "12kg"
    t.cell(2, 2).text = ""
    p = tmp_path / "t.docx"
    d.save(p)
    els = parse_docx(p, SID, RouteId.R5, "r")
    kv = [e for e in els if e.element_type.value == "key_value"]
    rows = [e for e in els if e.element_type.value == "table_row"]
    assert kv and kv[0].content == "封面标题表" and kv[0].metadata["merged_title"]
    assert rows[0].metadata["header"] == ["参数", "值", "备注"]  # 表头取下一行
    assert rows[0].content == "质量 | 12kg | "


def test_empty_table_rows_not_emitted(sample_docx):
    d = docx.Document(str(sample_docx))
    d.tables[0].rows[1].cells[0].text = ""
    d.tables[0].rows[1].cells[1].text = ""  # 整行空 → 不产出
    p2 = sample_docx.parent / "s2.docx"
    d.save(p2)
    els = parse_docx(p2, SID, RouteId.R4, "standard")
    assert not [e for e in els if e.element_type.value == "table_row"]
