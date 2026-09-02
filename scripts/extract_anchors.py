#!/usr/bin/env python3
"""TASK-P00-004 补充: 问题池证据锚点提取（可复现核验）.

按 source_id 或相对路径提取结构锚点: 标题/编号节、表头、信号句（状态/版本/引用词
所在段落, 截断显示）。用于人工核验 question_pool.csv 的 evidence_locations 是否属实。

用法:
  python scripts/extract_anchors.py SRC-321374652A          # 按 source_id
  python scripts/extract_anchors.py --path "3d打印 归档/X.docx"
  python scripts/extract_anchors.py --all > anchors.txt     # 全部（重定向自查）

输出仅在本地终端, 不落盘、不入库（DATA_GOVERNANCE: 正文不写日志/报告）。
"""

import csv
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from profile_sample import ROOT, _normalize_docx, _doc_text_via_libreoffice  # noqa: E402

SIG = re.compile(r"(决定|待办|待确认|建议|通过|不通过|问题|结论|遗留|风险|版本|修订|引用|依据|参考文件)")


def load_manifest():
    with open("data_manifest.csv", encoding="utf-8") as f:
        return {r["source_id"]: r["relative_path"] for r in csv.DictReader(f)}


def anchors_docx(path: Path):
    import docx

    d = docx.Document(str(path))
    heads = [p.text.strip() for p in d.paragraphs
             if p.style.name.lower().startswith(("heading", "标题")) and p.text.strip()]
    print(f"[标题] ({len(heads)})")
    for h in heads:
        print("  ", h)
    print(f"[表格] {len(d.tables)} 个")
    for i, t in enumerate(d.tables):
        hdr = [c.text.strip().replace("\n", "")[:14] for c in t.rows[0].cells][:8]
        print(f"  表{i}: {len(t.rows)}行x{len(t.columns)}列 表头={hdr}")
    print(f"[信号句]（状态/版本/引用词所在段落, 截断60字）")
    for p in d.paragraphs:
        t = p.text.strip()
        if SIG.search(t) and 8 < len(t) < 90:
            print("  ", t[:60])


def anchors_doc(path: Path):
    text = _doc_text_via_libreoffice(path)
    print(f"[文本] {len(text)} 字符")
    for line in text.splitlines():
        line = line.strip()
        if SIG.search(line) and 8 < len(line) < 90:
            print("  ", line[:60])


def anchors_pdf(path: Path):
    import fitz

    doc = fitz.open(path)
    print(f"[页数] {len(doc)}  [书签] {len(doc.get_toc())}")
    for lvl, title, page in doc.get_toc():
        print("  " * lvl, title)
    chars = [len(doc[i].get_text()) for i in range(len(doc))]
    print(f"[每页字符] {chars}")
    doc.close()


def anchors_xlsx(path: Path):
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    print(f"[sheet] {wb.sheetnames}")
    for ws in wb.worksheets:
        print(f"  == {ws.title} ({ws.max_row}x{ws.max_column})")
        for row in ws.iter_rows(max_row=4, values_only=True):
            cells = [str(c)[:16] for c in row if c is not None][:8]
            if cells:
                print("   ", cells)
    wb.close()


def main(argv) -> int:
    manifest = load_manifest()
    if argv[:1] == ["--all"]:
        targets = list(manifest.items())
    elif argv[:1] == ["--path"]:
        targets = [("", argv[1])]
    else:
        sid = argv[0]
        if sid not in manifest:
            print(f"source_id 不在 manifest: {sid}", file=sys.stderr)
            return 1
        targets = [(sid, manifest[sid])]

    for sid, rel in targets:
        path = _normalize_docx(ROOT / rel)
        print("=" * 30, sid, rel)
        ext = path.suffix.lower()
        if ext == ".docx":
            anchors_docx(path)
        elif ext == ".doc":
            anchors_doc(path)
        elif ext == ".pdf":
            anchors_pdf(path)
        else:
            anchors_xlsx(path)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
